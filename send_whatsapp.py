# ============ EDIT THIS VALUE TO CHANGE THE MESSAGE ============
OFFER_LINE = "enjoy 20% discount on any service"   # {{2}} - offer line ({{1}} is the customer's name from Excel)
# ===================================================================

# Send birthday and anniversary WhatsApp templates via the Meta WhatsApp Cloud API
# to everyone in the Excel file whose birthday / anniversary is today.
#
# Usage:
#     python send_whatsapp.py                      send today's messages
#     python send_whatsapp.py --dry-run            show who would get what, send nothing
#     python send_whatsapp.py --dry-run --date YYYY-MM-DD
#                                                  same, pretending today is that date
#     python send_whatsapp.py --force              run again even though everyone due today is already sent
#     python send_whatsapp.py --wait 30            wait 30 s (instead of DELIVERY_WAIT_SECONDS) for delivery reports
#     python send_whatsapp.py --setup-webhook      one-time: open the webhook so Meta can verify it
#
# Credentials come from .env (WHATSAPP_ACCESS_TOKEN, WHATSAPP_PHONE_NUMBER_ID). Delivery reports
# also need NGROK_AUTHTOKEN, NGROK_DOMAIN, WHATSAPP_APP_SECRET, WHATSAPP_WEBHOOK_VERIFY_TOKEN and
# WHATSAPP_WABA_ID (see README, "Delivery reports").

import argparse
import calendar
import hashlib
import hmac
import json
import logging
import os
import re
import shutil
import subprocess
import sys
import threading
import time
from contextlib import contextmanager
from datetime import date, datetime, timedelta
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from logging.handlers import RotatingFileHandler
from pathlib import Path
from urllib.parse import parse_qs, urlparse

if os.name == "nt":
    import msvcrt
else:
    import fcntl

import requests
from dotenv import load_dotenv

GRAPH_API_VERSION = "v20.0"
DELIVERY_WAIT_SECONDS = 90  # after sending, how long to wait for Meta's delivery reports (--wait overrides)
REQUEST_TIMEOUT = 30  # seconds
LANGUAGE_CODE = "en"  # must match the templates' language in WhatsApp Manager

TEMPLATES = {
    "birthday": "birthday_wish_discount",
    "anniversary": "anniversary_wish_discount",
}

HERE = Path(__file__).resolve().parent
EXCEL_FILE = HERE / "Salon Customers Database.xlsx"  # every sheet is read
SENT_LOG = HERE / "sent_log.json"                    # today's sends, to avoid repeats
SENT_LOG_JS = HERE / "sent_log_data.js"              # same data for dashboard.html (works from file://)
SENT_LOG_LOCK = HERE / "sent_log.lock"               # held while sent_log.json is rewritten
RUN_SUMMARY_JS = HERE / "run_summary_data.js"        # last real run's SUMMARY, for dashboard.html
LOCK_FILE = HERE / "send_whatsapp.lock"              # stops two runs sending at once
WEBHOOK_LOG = HERE / "webhook.log"                   # delivery reports received, rejected requests
WEBHOOK_EVENTS = HERE / "webhook_events.jsonl"       # every report exactly as Meta sent it
NGROK_LOG = HERE / "ngrok.log"                       # the tunnel's own log (for when it won't start)
DEFAULT_COUNTRY_CODE = "91"                          # Excel numbers are stored without it

# A message's status in sent_log.json only moves forward: "pending", then "sent" when WhatsApp
# accepts it, then what Meta's delivery report says. "unconfirmed" is set by this script when a
# message is still "sent" after MAX_ATTEMPTS tries; a late delivery report can still move it on.
STATUS_RANK = {"pending": 0, "sent": 1, "unconfirmed": 1, "delivered": 2, "read": 3, "failed": 4}
DELIVERY_FIELDS = ("status", "status_at", "error_code", "error_title", "error_details", "error_hint")
DONE_STATUSES = ("delivered", "read")                   # confirmed on the phone: never sent again
MAX_ATTEMPTS = 2                                        # per person, per occasion, per day
RETRY_AFTER = timedelta(hours=2)                        # a "sent" message gets this long for its report

# Meta error codes worth explaining in plain language.
KNOWN_ERRORS = {
    190: "Access token is invalid or expired. Generate a new token and update WHATSAPP_ACCESS_TOKEN in .env.",
    100: "Invalid parameter. Check WHATSAPP_PHONE_NUMBER_ID in .env and the phone number in the Excel file.",
    131030: "Recipient number is not in the allowed list. In test mode, add it under 'To' in the Meta app's WhatsApp > API Setup page.",
    131026: "Message undeliverable. The recipient may not have WhatsApp or has not accepted the latest terms.",
    131049: "This recipient has hit Meta's marketing message limit (shared across all businesses, not just yours). Wait a day or more before retrying - this isn't a code or account problem.",
    131047: "More than 24 hours have passed since the recipient last replied. Only approved templates can be sent.",
    131031: "The WhatsApp Business account is locked or restricted.",
    133010: "The sender phone number is not registered with the Cloud API.",
    130429: "Rate limit hit. Wait a moment and try again.",
    80007: "Rate limit hit for this WhatsApp Business account. Wait and try again.",
    132000: "Number of template variables doesn't match the template (expected {{1}} name and {{2}} offer line).",
    132001: "Template not found or not approved yet. Check its status in WhatsApp Manager (it must be Active).",
    132005: "Template text is too long after filling in the variables. Shorten OFFER_LINE or the customer's name.",
    132007: "Template content violates WhatsApp policy.",
    132012: "Template variable format doesn't match what the template expects.",
    132015: "Template is paused due to low quality. Check its status in WhatsApp Manager.",
    132016: "Template is disabled. Check its status in WhatsApp Manager.",
}


class RequestFailed(Exception):
    """The request never got an HTTP response (timeout, no connection, ...).

    maybe_sent is False only when no connection was ever made, so WhatsApp certainly didn't get it.
    """

    def __init__(self, message, maybe_sent):
        super().__init__(message)
        self.maybe_sent = maybe_sent


def never_connected(exc):
    """True if requests failed before any data reached the server (DNS failure, refused, connect timeout)."""
    from urllib3.exceptions import ConnectTimeoutError, NewConnectionError

    if isinstance(exc, requests.exceptions.ConnectTimeout):
        return True
    reason = getattr(exc.args[0], "reason", None) if exc.args else None
    return isinstance(reason, (NewConnectionError, ConnectTimeoutError))


def load_config():
    """Load credentials from the .env file next to this script."""
    load_dotenv(HERE / ".env")

    config = {
        "token": os.getenv("WHATSAPP_ACCESS_TOKEN", "").strip(),
        "phone_number_id": os.getenv("WHATSAPP_PHONE_NUMBER_ID", "").strip(),
    }
    missing = [
        name
        for name, key in (("WHATSAPP_ACCESS_TOKEN", "token"), ("WHATSAPP_PHONE_NUMBER_ID", "phone_number_id"))
        if not config[key]
    ]
    if missing:
        sys.exit(f"Error: missing required value(s) in .env: {', '.join(missing)}. See .env.example.")

    return config


def build_payload(to, template_name, variables):
    return {
        "messaging_product": "whatsapp",
        "to": to,
        "type": "template",
        "template": {
            "name": template_name,
            "language": {"code": LANGUAGE_CODE},
            "components": [
                {
                    "type": "body",
                    "parameters": [{"type": "text", "text": str(v)} for v in variables],
                }
            ],
        },
    }


def send_message(token, phone_number_id, payload):
    """POST the payload and return the requests.Response."""
    url = f"https://graph.facebook.com/{GRAPH_API_VERSION}/{phone_number_id}/messages"
    headers = {
        "Authorization": f"Bearer {token}",
        "Content-Type": "application/json",
    }
    return requests.post(url, headers=headers, json=payload, timeout=REQUEST_TIMEOUT)


def send_template(config, recipient, template_name, variables):
    """Send one template message. Returns (response, parsed body); raises RequestFailed on network errors."""
    payload = build_payload(recipient, template_name, variables)
    try:
        response = send_message(config["token"], config["phone_number_id"], payload)
    except requests.exceptions.Timeout as exc:
        raise RequestFailed(f"request timed out after {REQUEST_TIMEOUT}s. Check your connection and try again.",
                            maybe_sent=not never_connected(exc))
    except requests.exceptions.ConnectionError as exc:
        raise RequestFailed("could not reach graph.facebook.com. Check your internet connection.",
                            maybe_sent=not never_connected(exc))
    except requests.exceptions.RequestException as exc:
        raise RequestFailed(f"request failed: {exc}", maybe_sent=not never_connected(exc))

    try:
        body = response.json()
    except ValueError:
        body = {"raw": response.text}
    return response, body


def message_id_of(body):
    return (body.get("messages") or [{}])[0].get("id", "n/a")


def explain_error(body):
    """Build a readable message from a Graph API error response."""
    error = body.get("error", {}) if isinstance(body, dict) else {}
    code = error.get("code")
    details = error.get("error_data", {}).get("details") or error.get("message", "Unknown error")
    hint = KNOWN_ERRORS.get(code)
    lines = [f"API error {code}: {details}"]
    if hint:
        lines.append(f"Hint: {hint}")
    return "\n".join(lines)


# ------------------------------------------------------------------ Excel file

MONTHS = {name.lower(): i for i, name in enumerate(calendar.month_name) if name}
MONTHS.update({name.lower(): i for i, name in enumerate(calendar.month_abbr) if name})
MONTHS["sept"] = 9

REQUIRED_COLUMNS = ("First_Name", "Mobile_1", "Birth_Day", "Birth_Month", "Anniversary_Day", "Anniversary_Month")


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def parse_day_month(day, month):
    """Return (day, month_number), or None if either cell is blank or unreadable."""
    day_text, month_text = cell_text(day), cell_text(month).lower().rstrip(".")
    if not day_text or not month_text:
        return None
    if not day_text.isdigit() or month_text not in MONTHS:
        raise ValueError(f"unreadable date {day_text!r} {cell_text(month)!r}")
    return int(day_text), MONTHS[month_text]


def excel_phone(value):
    """Turn an Excel phone cell into API digits (country code added), or '' if blank/invalid."""
    digits = re.sub(r"\D", "", cell_text(value))
    if len(digits) == 11 and digits.startswith("0"):
        digits = digits[1:]
    if len(digits) == 10:
        digits = DEFAULT_COUNTRY_CODE + digits
    return digits if 11 <= len(digits) <= 15 else ""


def load_customers(path):
    """Yield one dict per data row from every sheet, looking columns up by header name."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            header = [cell_text(h) for h in next(rows, [])]
            missing = [c for c in REQUIRED_COLUMNS if c not in header]
            if missing:
                print(f"WARNING: sheet '{sheet.title}' skipped - missing column(s): {', '.join(missing)}")
                continue
            col = {name: header.index(name) for name in header if name}

            def get(row, name):
                i = col.get(name)
                return row[i] if i is not None and i < len(row) else None

            for row_number, row in enumerate(rows, start=2):
                if not any(cell_text(v) for v in row):
                    continue
                yield {
                    "where": f"{sheet.title} row {row_number}",
                    "name": " ".join(p for p in (cell_text(get(row, "First_Name")), cell_text(get(row, "Last_Name"))) if p),
                    "mobile_1": get(row, "Mobile_1"),
                    "mobile_2": get(row, "Mobile_2"),
                    "birthday": (get(row, "Birth_Day"), get(row, "Birth_Month")),
                    "anniversary": (get(row, "Anniversary_Day"), get(row, "Anniversary_Month")),
                }
    finally:
        workbook.close()


def find_todays_sends(customers, today):
    """Return one entry per (person, occasion) matching today; a person can appear twice."""
    sends, warnings = [], []
    for person in customers:
        for occasion in TEMPLATES:  # birthday, anniversary - checked independently
            try:
                day_month = parse_day_month(*person[occasion])
            except ValueError as exc:
                warnings.append(f"{person['where']} ({person['name']}): {occasion} {exc}, ignored")
                continue
            if day_month == (today.day, today.month):
                phone = excel_phone(person["mobile_1"]) or excel_phone(person["mobile_2"])
                raw = [cell_text(person[k]) for k in ("mobile_1", "mobile_2") if cell_text(person[k])]
                no_phone = f"invalid phone number ({', '.join(raw)})" if raw else "no phone number"
                sends.append({**person, "occasion": occasion, "phone": phone, "no_phone": no_phone})
    return sends, warnings


# ------------------------------------------------------- duplicate protection
#
# sent_log.json holds one entry per person and occasion sent to today (entries from earlier
# dates are removed at the start of each run). Delivery reports that arrive while the run
# waits update "sent" entries to delivered / read / failed.
# An entry is written as "pending" BEFORE the API call and becomes "sent" after it succeeds
# (or goes back to how it was if the API refuses it). What a later run does with an entry is
# decided by next_step(): only a "sent" one is ever resent, at most MAX_ATTEMPTS tries and
# RETRY_AFTER apart, overwriting the same entry (attempts counts the tries; earlier_message_ids
# keeps the old ids so a late delivery report still counts). The lock file stops two runs
# from sending at the same time.


def load_sent_log():
    if not SENT_LOG.exists():
        return []
    try:
        entries = json.loads(SENT_LOG.read_text(encoding="utf-8"))
    except (ValueError, OSError) as exc:
        sys.exit(f"Error: cannot read {SENT_LOG.name} ({exc}). Fix or move it; refusing to run without the duplicate check.")
    if not isinstance(entries, list):
        sys.exit(f"Error: {SENT_LOG.name} should contain a JSON list. Fix or move it.")
    return entries


def write_atomically(path, text):
    tmp = path.with_suffix(".tmp")
    tmp.write_text(text, encoding="utf-8")
    for attempt in range(20):
        try:
            tmp.replace(path)
            return
        except PermissionError:  # Windows: a reader (e.g. the dashboard) has the file open for a moment
            if attempt == 19:
                raise
            time.sleep(0.1)


@contextmanager
def sent_log_locked():
    """Hold sent_log.lock, so the sending loop and the webhook thread never rewrite sent_log.json at once.

    The OS drops the lock if the process dies, so it can't be left behind.
    """
    with open(SENT_LOG_LOCK, "a+b") as f:
        if os.name == "nt":
            deadline = time.monotonic() + 60
            while True:
                try:
                    f.seek(0)
                    msvcrt.locking(f.fileno(), msvcrt.LK_NBLCK, 1)
                    break
                except OSError:
                    if time.monotonic() > deadline:
                        raise
                    time.sleep(0.05)
        else:
            fcntl.flock(f, fcntl.LOCK_EX)
        try:
            yield
        finally:
            if os.name == "nt":
                f.seek(0)
                msvcrt.locking(f.fileno(), msvcrt.LK_UNLCK, 1)


def write_sent_log(entries):
    """Write sent_log.json and its dashboard copy. Call with sent_log_locked() held."""
    data = json.dumps(entries, indent=2)
    write_atomically(SENT_LOG, data)
    # A copy for dashboard.html, loaded with <script src> because browsers block fetch() on file:// pages.
    # "window." rather than "const" so the dashboard can load it again on Refresh.
    write_atomically(SENT_LOG_JS, f"window.sentLogData = {data};\n")


def save_sent_log(entries):
    """Save the sending loop's entries, keeping any newer delivery status the webhook has written meanwhile."""
    with sent_log_locked():
        on_disk = {e["message_id"]: e for e in load_sent_log() if isinstance(e, dict) and e.get("message_id")}
        for entry in entries:
            newer = on_disk.get(entry.get("message_id"))
            if newer and STATUS_RANK.get(newer.get("status"), 0) > STATUS_RANK.get(entry.get("status"), 0):
                entry.update({k: newer[k] for k in DELIVERY_FIELDS if k in newer})
        write_sent_log(entries)


def save_run_summary(today_str, counts, skipped_no_phone, skipped_duplicate, failures, waiting, needs_check):
    def flagged(items):
        return [{"name": s["name"], "occasion": s["occasion"], "phone": s["phone"], "status": e.get("status"),
                 "error_code": e.get("error_code"), "reason": why} for s, e, why in items]

    summary = {
        "date": today_str,
        "finished_at": datetime.now().isoformat(timespec="seconds"),
        "birthday_sent": counts["birthday"],
        "anniversary_sent": counts["anniversary"],
        "skipped_no_phone": [{"name": s["name"], "occasion": s["occasion"], "where": s["where"],
                              "reason": s["no_phone"]} for s in skipped_no_phone],
        "skipped_duplicate": [{"name": s["name"], "occasion": s["occasion"], "phone": s["phone"]}
                              for s in skipped_duplicate],
        "failed": [{"name": s["name"], "occasion": s["occasion"], "phone": s["phone"], "reason": reason,
                    "error_code": code} for s, reason, code in failures],
        "waiting": flagged(waiting),
        "needs_check": flagged(needs_check),
    }
    write_atomically(RUN_SUMMARY_JS, f"window.runSummaryData = {json.dumps(summary, indent=2)};\n")


# ------------------------------------------------------------ already processed today?
#
# Every run reads today's list from the Excel file first, then looks at sent_log.json:
#   - no entries dated today   -> a new day: full run
#   - some entries dated today -> compare Excel with them, names then mobile numbers, and
#     decide per person with next_step(). Nothing to send or mark -> "Nothing to do".
#
# next_step() is the one place that decides, from today's log entry for a mobile number and
# occasion, what a run does:
#   no entry                          -> send
#   delivered / read                  -> done, skip
#   sent, last try < RETRY_AFTER ago  -> wait for its delivery report
#   sent, MAX_ATTEMPTS tries made     -> mark "unconfirmed", never retried
#   sent, fewer tries                 -> resend (overwrites the same entry)
#   pending / failed / unconfirmed    -> flag for a manual check, never resent automatically


def todays_sent_log(today):
    return [e for e in load_sent_log() if isinstance(e, dict) and e.get("date") == today.isoformat()]


def clock(iso):
    try:
        return datetime.fromisoformat(iso).strftime("%H:%M")
    except (TypeError, ValueError):
        return "?"


def next_step(entry, now):
    """Return (action, why): action is "send", "resend", "done", "wait", "give_up" or "check"."""
    if entry is None:
        return "send", ""
    status = entry.get("status")
    attempts = entry.get("attempts", 1)
    if status in DONE_STATUSES:
        return "done", f"already {status} today"
    if status == "sent":
        try:
            last_try = datetime.fromisoformat(entry.get("sent_at"))
        except (TypeError, ValueError):
            last_try = None
        if last_try and now - last_try < RETRY_AFTER:
            return "wait", (f"sent at {clock(entry.get('sent_at'))} (try {attempts} of {MAX_ATTEMPTS}), "
                            f"no delivery report yet; not retried before {(last_try + RETRY_AFTER):%H:%M}")
        if attempts >= MAX_ATTEMPTS:
            return "give_up", f"not confirmed after {attempts} tries, check manually"
        return "resend", f"no delivery report since {clock(entry.get('sent_at'))} (try {attempts + 1} of {MAX_ATTEMPTS})"
    if status == "pending":
        return "check", "a run stopped while sending it, so it may have gone out; not resent automatically, check manually"
    if status == "failed":
        code = entry.get("error_code")
        return "check", (f"Meta reported it failed{f' (error {code})' if code else ''}"
                         f"{': ' + KNOWN_ERRORS[code].rstrip('.') if code in KNOWN_ERRORS else ''}; "
                         f"not resent automatically, check manually")
    if status == "unconfirmed":
        return "check", f"not confirmed after {attempts} tries, check manually"
    return "check", f"unknown status {status!r} in {SENT_LOG.name}; check manually"


def compare_with_sent_log(sends, todays_log, now):
    """For a day already processed: compare today's Excel list with today's sent_log.json entries.
    Returns (names_missing, to_send, to_mark). names_missing = people whose name isn't in the log
    at all; to_send = people next_step() would send to (new or changed number, or a "sent" due
    for its retry); to_mark = people about to be marked "unconfirmed". Rows without a phone or
    name are left out."""
    sends = [s for s in sends if s["phone"] and s["name"]]
    logged_names = {(str(e.get("name", "")).casefold(), e.get("occasion")) for e in todays_log}
    by_number = {(e.get("phone"), e.get("occasion")): e for e in todays_log}
    names_missing = [s for s in sends if (s["name"].casefold(), s["occasion"]) not in logged_names]
    steps = [(s, next_step(by_number.get((s["phone"], s["occasion"])), now)[0]) for s in sends]
    to_send = [s for s, action in steps if action in ("send", "resend")]
    to_mark = [s for s, action in steps if action == "give_up"]
    return names_missing, to_send, to_mark


def describe_new(new):
    return ", ".join(f"{s['name']} ({s['occasion']})" for s in new)


def acquire_lock():
    try:
        fd = os.open(LOCK_FILE, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
    except FileExistsError:
        started = LOCK_FILE.read_text(encoding="utf-8", errors="replace").strip() or "unknown time"
        sys.exit(
            f"Error: another run is already sending (started {started}).\n"
            f"If no other run is active (e.g. the last one was killed), delete {LOCK_FILE.name} and try again."
        )
    with os.fdopen(fd, "w", encoding="utf-8") as f:
        f.write(f"{datetime.now().isoformat(timespec='seconds')} by process {os.getpid()}")


def release_lock():
    try:
        LOCK_FILE.unlink()
    except FileNotFoundError:
        pass


# ---------------------------------------------------------- delivery reports
#
# Meta reports what happened to each message (delivered / read / failed) by calling a webhook
# URL. During a real run this script starts a small web server on this PC plus an ngrok tunnel
# that gives it the fixed public address https://NGROK_DOMAIN/webhook (registered once in the
# Meta App Dashboard). After sending it waits for the reports, records them in sent_log.json,
# and stops both. Nothing keeps running once the script ends.

HEALTH_TEXT = "WhatsApp webhook receiver is running."
MATCH_WAIT_SECONDS = 120  # a report can arrive before the sending loop has saved the message id
MAX_BODY_BYTES = 1_000_000
DELIVERY_SETTINGS = ("NGROK_AUTHTOKEN", "NGROK_DOMAIN", "WHATSAPP_APP_SECRET",
                     "WHATSAPP_WEBHOOK_VERIFY_TOKEN", "WHATSAPP_WABA_ID")

webhook_log = logging.getLogger("webhook")
webhook_events = logging.getLogger("webhook.events")


class DeliveryReportsUnavailable(Exception):
    pass


def setup_webhook_logging():
    if webhook_log.handlers:
        return
    handler = RotatingFileHandler(WEBHOOK_LOG, maxBytes=1_000_000, backupCount=3, encoding="utf-8")
    handler.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(message)s", "%Y-%m-%d %H:%M:%S"))
    webhook_log.addHandler(handler)
    webhook_log.setLevel(logging.INFO)
    webhook_log.propagate = False
    events_handler = RotatingFileHandler(WEBHOOK_EVENTS, maxBytes=5_000_000, backupCount=3, encoding="utf-8")
    events_handler.setFormatter(logging.Formatter("%(message)s"))
    webhook_events.addHandler(events_handler)
    webhook_events.setLevel(logging.INFO)
    webhook_events.propagate = False


def local_time(timestamp):
    try:
        return datetime.fromtimestamp(int(timestamp)).isoformat(timespec="seconds")
    except (TypeError, ValueError):
        return datetime.now().isoformat(timespec="seconds")


def apply_status(status):
    """Record one delivery report in its sent_log.json entry. False if no entry has that message id (yet)."""
    message_id, new = status.get("id"), status.get("status")
    with sent_log_locked():
        entries = load_sent_log()
        entry = next((e for e in entries if isinstance(e, dict) and e.get("message_id") == message_id), None)
        if entry is None:
            # A late report for an earlier attempt of a message that was sent again. Only a
            # delivered/read counts (that person did get it); anything else is about an old try.
            entry = next((e for e in entries if isinstance(e, dict)
                          and message_id in e.get("earlier_message_ids", ())), None)
            if entry is None:
                return False
            if new not in DONE_STATUSES:
                return True
        if STATUS_RANK[new] <= STATUS_RANK.get(entry.get("status"), 0):
            return True  # arrived out of order, e.g. "delivered" after "read": keep the later state
        entry["status"] = new
        entry["status_at"] = local_time(status.get("timestamp"))
        if new == "failed":
            error = (status.get("errors") or [{}])[0]
            code = error.get("code")
            entry["error_code"] = code
            entry["error_title"] = error.get("title") or error.get("message") or "Unknown error"
            entry["error_details"] = (error.get("error_data") or {}).get("details") or error.get("message") or ""
            entry["error_hint"] = KNOWN_ERRORS.get(code, "")
        write_sent_log(entries)
    webhook_log.info(f"{new} {entry.get('name')} +{entry.get('phone')} ({entry.get('occasion')})"
                     + (f": error {entry['error_code']} {entry['error_title']}" if new == "failed" else ""))
    return True


def handle_status(status):
    if status.get("status") not in STATUS_RANK:
        return  # a kind of report the dashboard doesn't track
    deadline = time.monotonic() + MATCH_WAIT_SECONDS
    try:
        while not apply_status(status):
            if time.monotonic() > deadline:
                # Usually a late report for an earlier day's message (the log only keeps today's).
                webhook_log.info(f"no {SENT_LOG.name} entry for {status.get('status')} report of {status.get('id')}; ignored")
                return
            time.sleep(1)
    except SystemExit as exc:  # load_sent_log() exits on an unreadable sent_log.json
        webhook_log.error(f"could not record report for {status.get('id')}: {exc}")
    except Exception:
        webhook_log.exception(f"could not record report for {status.get('id')}")


def process_webhook(payload):
    for entry in payload.get("entry") or []:
        for change in entry.get("changes") or []:
            value = change.get("value") or {}
            for status in value.get("statuses") or []:
                webhook_events.info(json.dumps({"received_at": datetime.now().isoformat(timespec="seconds"), **status}))
                handle_status(status)
            for message in value.get("messages") or []:
                webhook_log.info(f"incoming message from +{message.get('from')} ({message.get('type')})")


class WebhookHandler(BaseHTTPRequestHandler):
    server_version = "WhatsAppWebhookReceiver"
    verify_token = b""   # set per server in DeliveryReports.start()
    app_secret = b""
    verified = None      # threading.Event, set when Meta verifies the callback URL

    def do_GET(self):
        url = urlparse(self.path)
        if url.path == "/":
            return self.reply(200, HEALTH_TEXT)
        if url.path != "/webhook":
            return self.reply(404, "Not found")
        # Meta's check when the callback URL is saved in the App Dashboard.
        query = {k: v[0] for k, v in parse_qs(url.query).items()}
        token = query.get("hub.verify_token", "").encode()
        if query.get("hub.mode") == "subscribe" and hmac.compare_digest(token, self.verify_token):
            webhook_log.info("Meta verified the callback URL")
            self.verified.set()
            return self.reply(200, query.get("hub.challenge", ""))
        webhook_log.warning("callback URL check with a wrong verify token was rejected")
        self.reply(403, "Verify token doesn't match")

    def do_POST(self):
        if urlparse(self.path).path != "/webhook":
            return self.reply(404, "Not found")
        try:
            length = int(self.headers.get("Content-Length") or 0)
        except ValueError:
            length = -1
        if not 0 <= length <= MAX_BODY_BYTES:
            return self.reply(413, "Bad length")
        body = self.rfile.read(length)
        # Only Meta knows the app secret, so a valid signature proves the report came from Meta.
        expected = "sha256=" + hmac.new(self.app_secret, body, hashlib.sha256).hexdigest()
        if not hmac.compare_digest(self.headers.get("X-Hub-Signature-256", "").encode(), expected.encode()):
            webhook_log.warning("POST with a missing or wrong signature was rejected")
            return self.reply(403, "Bad signature")
        self.reply(200, "OK")  # answer Meta straight away; recording may wait for the sending loop
        try:
            payload = json.loads(body)
        except ValueError:
            webhook_log.warning("signed POST was not valid JSON; ignored")
            return
        threading.Thread(target=process_webhook, args=(payload,), daemon=True).start()

    def reply(self, code, text):
        data = text.encode()
        self.send_response(code)
        self.send_header("Content-Type", "text/plain; charset=utf-8")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def log_message(self, format, *args):
        pass  # keep the run's output clean; the useful parts go to webhook.log


def graph_request(method, path):
    token = os.getenv("WHATSAPP_ACCESS_TOKEN", "").strip()
    url = f"https://graph.facebook.com/{GRAPH_API_VERSION}/{path}"
    response = requests.request(method, url, headers={"Authorization": f"Bearer {token}"}, timeout=REQUEST_TIMEOUT)
    try:
        return response.ok, response.json()
    except ValueError:
        return response.ok, {"raw": response.text}


def ensure_waba_subscription(waba_id):
    """Meta sends no reports unless the app is subscribed to the WhatsApp Business Account. Idempotent."""
    ok, body = graph_request("GET", f"{waba_id}/subscribed_apps")
    if ok and body.get("data"):
        return
    ok, body = graph_request("POST", f"{waba_id}/subscribed_apps")
    if not ok:
        raise DeliveryReportsUnavailable(f"could not subscribe the app to WhatsApp Business Account {waba_id}: "
                                         f"{explain_error(body)}")
    print("Delivery reports: subscribed the app to the WhatsApp Business Account's webhooks (one-time).")


class DeliveryReports:
    """The webhook server and ngrok tunnel for one run. start() before sending, stop() when done."""

    def __init__(self, settings):
        self.settings = settings
        self.domain = settings["NGROK_DOMAIN"].removeprefix("https://").strip("/")
        self.url = f"https://{self.domain}/webhook"
        self.server = None
        self.ngrok = None
        self.verified = threading.Event()

    @classmethod
    def from_env(cls):
        """A DeliveryReports if .env has every setting it needs; otherwise None, after saying what's missing."""
        settings = {name: os.getenv(name, "").strip() for name in DELIVERY_SETTINGS}
        missing = [name for name, value in settings.items() if not value]
        if len(missing) == len(settings):
            print("Delivery reports: off (not set up; see README, \"Delivery reports\").\n")
            return None
        if missing:
            print(f"Delivery reports: off - missing in .env: {', '.join(missing)}\n")
            return None
        return cls(settings)

    @property
    def running(self):
        return self.ngrok is not None

    def start(self):
        """Start the local server and the tunnel, and check the public URL answers. Raises DeliveryReportsUnavailable."""
        if self.running:
            return
        setup_webhook_logging()
        handler = type("Handler", (WebhookHandler,), {
            "verify_token": self.settings["WHATSAPP_WEBHOOK_VERIFY_TOKEN"].encode(),
            "app_secret": self.settings["WHATSAPP_APP_SECRET"].encode(),
            "verified": self.verified,
        })
        self.server = ThreadingHTTPServer(("127.0.0.1", 0), handler)  # any free port; only ngrok connects to it
        threading.Thread(target=self.server.serve_forever, daemon=True).start()
        port = self.server.server_address[1]

        exe = os.getenv("NGROK_PATH", "").strip() or shutil.which("ngrok")
        if not exe:
            self.stop()
            raise DeliveryReportsUnavailable("ngrok is not installed or not on PATH "
                                             "(winget install ngrok.ngrok, then open a new terminal)")
        env = {**os.environ, "NGROK_AUTHTOKEN": self.settings["NGROK_AUTHTOKEN"]}
        self.ngrok = subprocess.Popen(
            [exe, "http", str(port), f"--url=https://{self.domain}", "--log", str(NGROK_LOG), "--log-format", "logfmt"],
            env=env, stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)

        # Ready once the public address reaches our server (proves ngrok, the domain and the server all work).
        deadline = time.monotonic() + 30
        while True:
            if self.ngrok.poll() is not None:
                problem = self.ngrok_error()
                self.stop()
                raise DeliveryReportsUnavailable(f"ngrok stopped: {problem}")
            try:
                r = requests.get(f"https://{self.domain}/", headers={"ngrok-skip-browser-warning": "1"}, timeout=5)
                if r.text == HEALTH_TEXT:
                    break
            except requests.RequestException:
                pass
            if time.monotonic() > deadline:
                problem = self.ngrok_error()
                self.stop()
                raise DeliveryReportsUnavailable(f"https://{self.domain}/ didn't answer within 30 s. {problem}")
            time.sleep(1)

        try:
            ensure_waba_subscription(self.settings["WHATSAPP_WABA_ID"])
        except (requests.RequestException, DeliveryReportsUnavailable) as exc:
            self.stop()
            raise DeliveryReportsUnavailable(str(exc))
        webhook_log.info(f"listening at {self.url}")

    def ngrok_error(self):
        try:
            lines = NGROK_LOG.read_text(encoding="utf-8", errors="replace").splitlines()
        except OSError:
            return "no ngrok.log"
        errors = [line for line in lines if "lvl=eror" in line or "lvl=crit" in line or "err=" in line]
        detail = (errors or lines or ["(ngrok.log is empty)"])[-1]
        hint = ""
        if "ERR_NGROK_334" in detail or "already online" in detail:
            hint = " Another ngrok is still using this domain: end ngrok.exe in Task Manager and run again."
        elif "ERR_NGROK_105" in detail or "authtoken" in detail.lower():
            hint = " Check NGROK_AUTHTOKEN in .env."
        return f"{detail}{hint}"

    def wait_for_reports(self, message_ids, seconds):
        """Show reports for this run's messages as they arrive; return when all have one or time is up."""
        if not message_ids or seconds <= 0:
            return
        print(f"\nWaiting up to {seconds} s for delivery reports (Ctrl+C to stop waiting)...")
        shown = {}
        deadline = time.monotonic() + seconds
        try:
            while True:
                entries = {e.get("message_id"): e for e in load_sent_log() if isinstance(e, dict)}
                for mid in message_ids:
                    e = entries.get(mid, {})
                    status = e.get("status")
                    if status != shown.get(mid) and STATUS_RANK.get(status, 0) >= STATUS_RANK["delivered"]:
                        who = f"{e.get('name')} +{e.get('phone')} ({e.get('occasion')})"
                        if status == "failed":
                            print(f"    FAILED    {who}: error {e.get('error_code')} {e.get('error_title')}")
                            if e.get("error_hint"):
                                print(f"              {e['error_hint']}")
                        else:
                            print(f"    {status.upper():9} {who}")
                    shown[mid] = status
                done = [m for m in message_ids if STATUS_RANK.get(shown.get(m), 0) >= STATUS_RANK["delivered"]]
                if len(done) == len(message_ids) or time.monotonic() > deadline:
                    break
                time.sleep(1)
        except KeyboardInterrupt:
            print("    (stopped waiting)")
        final = [shown.get(m) for m in message_ids]
        delivered = sum(s in ("delivered", "read") for s in final)
        failed = final.count("failed")
        waiting = len(final) - delivered - failed
        print(f"Delivery reports: {delivered} delivered, {failed} failed, {waiting} with no report yet")
        if waiting:
            print("    ('no report yet' stays 'sent' on the dashboard: the phone may be off, or Meta was slow.)")

    def stop(self):
        if self.ngrok is not None:
            self.ngrok.terminate()
            try:
                self.ngrok.wait(timeout=10)
            except subprocess.TimeoutExpired:
                self.ngrok.kill()
            self.ngrok = None
        if self.server is not None:
            self.server.shutdown()
            self.server.server_close()
            self.server = None


def setup_webhook():
    """One-time: open the webhook until Meta verifies the callback URL in the App Dashboard."""
    load_dotenv(HERE / ".env")
    reports = DeliveryReports.from_env()
    if reports is None:
        return 1
    try:
        reports.start()
    except DeliveryReportsUnavailable as exc:
        print(f"Error: {exc}")
        return 1
    try:
        print(f"The webhook is open. In the Meta App Dashboard > your app > WhatsApp > Configuration > Webhook > Edit:\n"
              f"    Callback URL:  {reports.url}\n"
              f"    Verify token:  {reports.settings['WHATSAPP_WEBHOOK_VERIFY_TOKEN']}\n"
              f"then click 'Verify and save'. Waiting up to 15 minutes (Ctrl+C to cancel)...")
        if not reports.verified.wait(timeout=15 * 60):
            print("No verification arrived. Check the URL and token, then run --setup-webhook again.")
            return 1
        print("Meta verified the callback URL. Last step in the same page: under 'Webhook fields', "
              "Subscribe to 'messages'. Setup is done; every run now collects delivery reports.")
        time.sleep(3)  # let Meta finish its request before the tunnel closes
        return 0
    except KeyboardInterrupt:
        return 1
    finally:
        reports.stop()


# ------------------------------------------------------------------------ run


def load_todays_sends(excel_path, today):
    """Read the Excel file: (messages due today, warnings)."""
    if not excel_path.exists():
        sys.exit(f"Error: Excel file not found: {excel_path}")
    return find_todays_sends(load_customers(excel_path), today)


def run(today, dry_run, excel_path, sends, warnings, reports=None, wait_seconds=DELIVERY_WAIT_SECONDS):
    if not OFFER_LINE.strip():
        sys.exit("Error: OFFER_LINE cannot be empty.")
    config = None if dry_run else load_config()

    today_str = today.isoformat()
    # The log only ever holds today's entries: earlier dates are dropped (and saved, on a real run)
    # before anything is sent. Today's own entries stay, so re-runs today still skip them.
    full_log = load_sent_log()
    sent_log = [e for e in full_log if isinstance(e, dict) and e.get("date") == today_str]
    removed_old = len(full_log) - len(sent_log)
    if removed_old and not dry_run:
        save_sent_log(sent_log)
    logged = {(e.get("date"), e.get("phone"), e.get("occasion")): e for e in sent_log}

    mode = "DRY RUN - nothing will be sent" if dry_run else "sending"
    print(f"Birthday/anniversary messages for {today.day} {calendar.month_name[today.month]} ({today_str}), {mode}")
    print(f"Excel: {excel_path.name}  ->  {len(sends)} message(s) due today")
    if removed_old:
        print(f"{SENT_LOG.name}: {removed_old} entr{'y' if removed_old == 1 else 'ies'} from earlier dates "
              + ("would be removed on a real run" if dry_run else "removed"))
    print()
    for w in warnings:
        print(f"WARNING: {w}")
    if warnings:
        print()

    counts = {"birthday": 0, "anniversary": 0}
    skipped_no_phone, skipped_duplicate, failures = [], [], []  # failures: (send, reason, error code)
    waiting, needs_check = [], []                                # (send, log entry, why)
    accepted = []  # (send, message id) for each message WhatsApp accepted this run
    tried_this_run = set()  # a person on two Excel rows still gets one message per run

    for s in sends:
        template_name = TEMPLATES[s["occasion"]]
        label = f"[{s['occasion'].upper():11}] {s['name'] or '(no name)'} | " \
                f"{'+' + s['phone'] if s['phone'] else 'no phone'} | {template_name} | {s['where']}"

        if not s["phone"]:
            print(f"{label} -> SKIPPED: {s['no_phone']}")
            skipped_no_phone.append(s)
            continue
        if not s["name"]:
            print(f"{label} -> FAILED: no name in First_Name")
            failures.append((s, "no name in First_Name", None))
            continue
        key = (today_str, s["phone"], s["occasion"])
        if key in tried_this_run:
            print(f"{label} -> SKIPPED: already handled in this run")
            continue
        tried_this_run.add(key)
        existing = logged.get(key)
        action, why = next_step(existing, datetime.now())
        if action == "done":
            print(f"{label} -> SKIPPED: {why}")
            skipped_duplicate.append(s)
            continue
        if action == "wait":
            print(f"{label} -> WAITING: {why}")
            waiting.append((s, existing, why))
            continue
        if action == "check":
            print(f"{label} -> CHECK MANUALLY: {why}")
            needs_check.append((s, existing, why))
            continue
        if action == "give_up":
            if dry_run:
                print(f"{label} -> WOULD MARK UNCONFIRMED: {why}")
            else:
                existing.update(status="unconfirmed", status_at=datetime.now().isoformat(timespec="seconds"))
                save_sent_log(sent_log)
                print(f"{label} -> UNCONFIRMED: {why}")
            needs_check.append((s, existing, why))
            continue
        # "send" (no entry yet) or "resend" (still "sent" after RETRY_AFTER, tries left).
        resend_note = f" (again: {why})" if action == "resend" else ""
        if dry_run:
            print(f"{label} -> WOULD SEND{resend_note}  {{{{1}}}}={s['name']!r}  {{{{2}}}}={OFFER_LINE!r}")
            counts[s["occasion"]] += 1
            continue

        if reports and not reports.running:
            # Opened just before the first message, so reports can't arrive before we listen.
            try:
                reports.start()
                print(f"Delivery reports: listening at {reports.url}\n")
            except DeliveryReportsUnavailable as exc:
                print(f"WARNING: delivery reports off for this run - {exc}\n")
                reports = None

        before = dict(existing) if existing else None  # to put back if this attempt never leaves
        if existing:
            entry = existing  # one entry per person and occasion: overwrite it, don't add another
            earlier = entry.get("earlier_message_ids", [])
            if entry.get("message_id"):
                earlier = earlier + [entry.pop("message_id")]  # a late report for it still counts
            for field in DELIVERY_FIELDS:
                entry.pop(field, None)
            entry.update(name=s["name"], status="pending", sent_at=datetime.now().isoformat(timespec="seconds"),
                         attempts=entry.get("attempts", 1) + 1, earlier_message_ids=earlier)
        else:
            entry = {"date": today_str, "phone": s["phone"], "occasion": s["occasion"], "name": s["name"],
                     "status": "pending", "sent_at": datetime.now().isoformat(timespec="seconds"), "attempts": 1}
            sent_log.append(entry)
            logged[key] = entry
        save_sent_log(sent_log)  # recorded before sending, so a crash leaves it at "pending"

        def undo():
            """This attempt never left this computer / was refused: back to how it was before."""
            if before is None:
                sent_log.remove(entry)
                del logged[key]
            else:
                entry.clear()
                entry.update(before)
            save_sent_log(sent_log)

        try:
            response, body = send_template(config, s["phone"], template_name, [s["name"], OFFER_LINE])
        except RequestFailed as exc:
            code = None
            if exc.maybe_sent:
                # The request may have reached WhatsApp without us seeing the reply: the entry
                # stays "pending", which is never resent automatically (it may have gone out).
                reason = f"{exc} It may have been delivered; it stays 'pending' in {SENT_LOG.name} " \
                         f"and won't be resent automatically - check manually."
            else:
                undo()
                reason = str(exc)
        else:
            if response.ok:
                entry.update(status="sent", message_id=message_id_of(body))
                save_sent_log(sent_log)
                print(f"{label} -> SENT{resend_note} (id: {entry['message_id']})")
                accepted.append((s, entry["message_id"]))  # counted after the delivery reports
                continue
            undo()  # WhatsApp refused it, so nothing went out
            reason = explain_error(body).replace("\n", " | ")
            code = (body.get("error") or {}).get("code") if isinstance(body, dict) else None
        print(f"{label} -> FAILED: {reason}")
        failures.append((s, reason, code))

    if not dry_run:
        # Wait for the delivery reports BEFORE tallying, so the summary shows each message's
        # final status: one Meta reports as failed (e.g. 131049) counts as failed, not sent.
        if reports and reports.running:
            reports.wait_for_reports([mid for _, mid in accepted], wait_seconds)
        final = {e.get("message_id"): e for e in load_sent_log() if isinstance(e, dict) and e.get("message_id")}
        for s, mid in accepted:
            e = final.get(mid, {})
            if e.get("status") == "failed":
                code = e.get("error_code")
                # The code is kept separately (the dashboard shows it as a label), not in the text.
                reason = f"Meta reported it failed after sending: {e.get('error_title') or 'no reason given'}"
                failures.append((s, reason, code))
            else:
                counts[s["occasion"]] += 1  # sent (no report yet), delivered or read

    verb = "would be sent" if dry_run else "sent"
    print("\n=========== SUMMARY ===========")
    print(f"Birthday messages {verb}:     {counts['birthday']}")
    print(f"Anniversary messages {verb}:  {counts['anniversary']}")
    print(f"Skipped (no phone number):  {len(skipped_no_phone)}")
    for s in skipped_no_phone:
        print(f"    - {s['name']} ({s['occasion']}, {s['where']}): {s['no_phone']}")
    print(f"Skipped (already delivered today): {len(skipped_duplicate)}")
    print(f"Waiting for delivery report:  {len(waiting)}")
    for s, _, why in waiting:
        print(f"    - {s['name']} +{s['phone']} ({s['occasion']}): {why}")
    print(f"Needs a manual check:       {len(needs_check)}")
    for s, _, why in needs_check:
        print(f"    - {s['name']} +{s['phone']} ({s['occasion']}): {why}")
    print(f"Failed:                     {len(failures)}")
    for s, reason, _ in failures:
        print(f"    - {s['name']} +{s['phone']} ({s['occasion']}): {reason}")
    if not dry_run and sum(counts.values()) and not (reports and reports.running):
        print("\nNote: 'sent' means WhatsApp accepted the message. Meta can still drop it later "
              "(e.g. error 131049); without delivery reports that doesn't show anywhere.")
    if not dry_run:
        save_run_summary(today_str, counts, skipped_no_phone, skipped_duplicate, failures, waiting, needs_check)
        if failures:
            print(f"\n{len(failures)} message(s) failed. Messages Meta rejected after sending are not "
                  f"retried automatically; see the dashboard.")
    return 1 if failures else 0


def main():
    parser = argparse.ArgumentParser(
        description="Send the birthday / anniversary WhatsApp template to everyone in the Excel file "
                    "whose day is today. Safe to re-run: anyone already sent today is skipped "
                    f"(tracked in {SENT_LOG.name}).",
    )
    parser.add_argument("--dry-run", action="store_true", help="list who would get which message; send nothing")
    parser.add_argument("--date", metavar="YYYY-MM-DD", help="with --dry-run: pretend today is this date")
    parser.add_argument("--excel", metavar="PATH", help=f"Excel file to read (default: {EXCEL_FILE.name})")
    parser.add_argument("--force", action="store_true",
                        help=f"run even if every number due today is already in {SENT_LOG.name} "
                             f"(each person is still handled as usual: delivered skipped, retry limits kept)")
    parser.add_argument("--wait", metavar="SECONDS", type=int, default=DELIVERY_WAIT_SECONDS,
                        help=f"after sending, wait up to this long for delivery reports (default {DELIVERY_WAIT_SECONDS}; 0 = don't wait)")
    parser.add_argument("--setup-webhook", action="store_true",
                        help="one-time setup: open the webhook so Meta can verify its callback URL, then stop")
    args = parser.parse_args()

    if args.setup_webhook:
        return setup_webhook()

    if args.date and not args.dry_run:
        parser.error("--date only works with --dry-run, so real messages always go out on the real date")
    try:
        today = date.fromisoformat(args.date) if args.date else date.today()
    except ValueError:
        parser.error(f"--date {args.date!r} is not a valid YYYY-MM-DD date")
    excel_path = Path(args.excel).resolve() if args.excel else EXCEL_FILE

    # Step 1: today's list from the Excel file, before any other check.
    sends, warnings = load_todays_sends(excel_path, today)

    def already_processed_check():
        """Steps 2-4. Returns the message to print when there's nothing to send, else None."""
        todays_log = todays_sent_log(today)
        if not todays_log:
            return None  # nothing sent today yet: a new day, full run
        names_missing, to_send, to_mark = compare_with_sent_log(sends, todays_log, datetime.now())
        if names_missing:
            print(f"{SENT_LOG.name} already has {len(todays_log)} message(s) today, but these names due "
                  f"today aren't in it: {describe_new(names_missing)}.")
        else:
            print(f"{SENT_LOG.name} already has {len(todays_log)} message(s) today and every name due "
                  f"today is in it; checking mobile numbers.")
        if not to_send and not to_mark:
            return (f"Nobody due today needs a message now (each is delivered, waiting for a delivery "
                    f"report, or flagged for a manual check on the dashboard). Nothing to do.\n"
                    f"To see each person's state, use --dry-run.")
        if to_send:
            print(f"To send (new or changed number, or a retry that's due): {describe_new(to_send)}.")
        if to_mark:
            print(f"To mark unconfirmed after {MAX_ATTEMPTS} tries: {describe_new(to_mark)}.")
        print("Everyone else is skipped.\n")
        return None

    if args.dry_run:
        stop = already_processed_check()
        if stop:
            print(f"(A real run would stop here.) {stop}\n")
        return run(today, True, excel_path, sends, warnings)
    acquire_lock()
    try:
        # Checked after taking the lock, so a run finishing right now can't be missed.
        stop = None if args.force else already_processed_check()
        if stop:
            print(stop)
            return 0
        load_dotenv(HERE / ".env")
        reports = DeliveryReports.from_env()
        try:
            return run(today, False, excel_path, sends, warnings, reports, args.wait)
        finally:
            if reports:
                reports.stop()  # nothing stays running after the script ends
    finally:
        release_lock()


if __name__ == "__main__":
    sys.exit(main())
